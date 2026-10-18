import html
import io
import re

import docx
import pymupdf
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.config import settings
from app.generators.render import date_range, join_parts
from app.llm.models import ContactInfo, CoverLetter, ParsedResume

# Single-column, standard fonts, real text (no images or tables) — what applicant tracking systems parse best.
PDF_CSS = """
body { font-family: sans-serif; font-size: 10pt; line-height: 1.3; }
h1 { font-size: 18pt; margin: 0 0 2pt 0; }
h2 { font-size: 11.5pt; margin: 10pt 0 3pt 0; border-bottom: 0.5pt solid #444; text-transform: uppercase; }
h3 { font-size: 10.5pt; margin: 6pt 0 1pt 0; }
p { margin: 0 0 3pt 0; }
ul { margin: 2pt 0 4pt 0; }
li { margin-bottom: 1pt; }
.meta { color: #444; }
"""

LETTER_CSS = "body { font-family: sans-serif; font-size: 11pt; line-height: 1.45; } p { margin: 0 0 10pt 0; }"


def safe_filename(*parts: str) -> str:
    name = "_".join(p for p in parts if p)
    return re.sub(r"[^A-Za-z0-9]+", "_", name).strip("_") or "document"


# --- PDF ---


def _esc(text: str) -> str:
    return html.escape(text, quote=False)


def _contact_lines(contact: ContactInfo) -> list[str]:
    return [
        join_parts(contact.email, contact.phone, contact.location),
        join_parts(contact.linkedin, contact.github, contact.portfolio, *contact.other_links),
    ]


def resume_html(resume: ParsedResume) -> str:
    parts = [f"<h1>{_esc(resume.contact.name)}</h1>"]
    parts += [f'<p class="meta">{_esc(line)}</p>' for line in _contact_lines(resume.contact) if line]

    def bullets(items: list[str]) -> str:
        return "<ul>" + "".join(f"<li>{_esc(i)}</li>" for i in items) + "</ul>" if items else ""

    if resume.summary:
        parts.append(f"<h2>Summary</h2><p>{_esc(resume.summary)}</p>")
    if resume.skills:
        parts.append(f"<h2>Skills</h2><p>{_esc(', '.join(resume.skills))}</p>")
    if resume.experience:
        parts.append("<h2>Experience</h2>")
        for e in resume.experience:
            parts.append(f"<h3>{_esc(e.title)} — {_esc(e.company)}</h3>")
            meta = join_parts(e.location, date_range(e.start_date, e.end_date))
            if meta:
                parts.append(f'<p class="meta">{_esc(meta)}</p>')
            parts.append(bullets(e.bullets))
    if resume.projects:
        parts.append("<h2>Projects</h2>")
        for p in resume.projects:
            tech = f" ({', '.join(p.technologies)})" if p.technologies else ""
            parts.append(f"<h3>{_esc(p.name)}{_esc(tech)}</h3>")
            meta = join_parts(p.description, p.link)
            if meta:
                parts.append(f'<p class="meta">{_esc(meta)}</p>')
            parts.append(bullets(p.bullets))
    if resume.education:
        parts.append("<h2>Education</h2>")
        for ed in resume.education:
            parts.append(f"<h3>{_esc(ed.institution)}</h3>")
            meta = join_parts(
                join_parts(ed.degree, ed.field_of_study, sep=", "), ed.location,
                date_range(ed.start_date, ed.end_date), ed.grade,
            )
            if meta:
                parts.append(f'<p class="meta">{_esc(meta)}</p>')
            parts.append(bullets(ed.highlights))
    for title, items in (
        ("Certifications", resume.certifications),
        ("Achievements", resume.achievements),
        ("Languages", resume.languages),
    ):
        if items:
            parts.append(f"<h2>{title}</h2>{bullets(items)}")
    return "".join(parts)


def cover_letter_html(letter: CoverLetter, contact: ContactInfo) -> str:
    header = [f"<p><b>{_esc(contact.name)}</b></p>"] if contact.name else []
    header += [f'<p>{_esc(line)}</p>' for line in _contact_lines(contact)[:1] if line]
    body = [f"<p>{_esc(letter.greeting)}</p>"] + [f"<p>{_esc(p)}</p>" for p in letter.paragraphs]
    closing = f"<p>{_esc(letter.sign_off)}<br/>{_esc(letter.signature)}</p>"
    return "".join(header) + "<br/>" + "".join(body) + closing


def _as_uri(text: str) -> str:
    if "@" in text and "/" not in text:
        return f"mailto:{text}"
    return text if text.startswith(("http://", "https://")) else f"https://{text}"


def link_targets(contact: ContactInfo, extra: list[str] | None = None) -> dict[str, str]:
    """Visible text -> URI for every link the document shows."""
    texts = [contact.email, contact.linkedin, contact.github, contact.portfolio, *contact.other_links]
    return {t: _as_uri(t) for t in [*texts, *(extra or [])] if t}


def _add_links(doc: pymupdf.Document, targets: dict[str, str]) -> None:
    # Longest text first, so "github.com/me/repo" isn't claimed by the shorter "github.com/me".
    for page_number in range(doc.page_count):
        page = doc[page_number]
        taken: list[pymupdf.Rect] = []
        for text in sorted(targets, key=len, reverse=True):
            for rect in page.search_for(text):
                if any(rect.intersects(t) for t in taken):
                    continue
                page.insert_link({"kind": pymupdf.LINK_URI, "from": rect, "uri": targets[text]})
                taken.append(rect)


def html_to_pdf(markup: str, css: str, links: dict[str, str] | None = None) -> bytes:
    story = pymupdf.Story(html=markup, user_css=css)
    page = pymupdf.paper_rect(settings.pdf_page_size)
    content_area = page + (50, 45, -50, -45)
    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)
    more = True
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(content_area)
        story.draw(device)
        writer.end_page()
    writer.close()

    doc = pymupdf.open(stream=buffer.getvalue(), filetype="pdf")
    if links:
        _add_links(doc, links)
    # Embed only the glyphs used, so a one-page resume isn't hundreds of KB.
    doc.subset_fonts()
    return doc.tobytes(garbage=3, deflate=True)


def resume_pdf(resume: ParsedResume) -> bytes:
    project_links = [p.link for p in resume.projects if p.link]
    return html_to_pdf(resume_html(resume), PDF_CSS, link_targets(resume.contact, project_links))


def cover_letter_pdf(letter: CoverLetter, contact: ContactInfo) -> bytes:
    return html_to_pdf(cover_letter_html(letter, contact), LETTER_CSS, link_targets(contact))


# --- DOCX ---


def _docx_bytes(document: docx.document.Document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _new_document(font_size: float) -> docx.document.Document:
    document = docx.Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(font_size)
    return document


def resume_docx(resume: ParsedResume) -> bytes:
    document = _new_document(10.5)
    name = document.add_heading(resume.contact.name, level=0)
    name.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for line in _contact_lines(resume.contact):
        if line:
            document.add_paragraph(line)

    def bullets(items: list[str]) -> None:
        for item in items:
            document.add_paragraph(item, style="List Bullet")

    if resume.summary:
        document.add_heading("Summary", level=1)
        document.add_paragraph(resume.summary)
    if resume.skills:
        document.add_heading("Skills", level=1)
        document.add_paragraph(", ".join(resume.skills))
    if resume.experience:
        document.add_heading("Experience", level=1)
        for e in resume.experience:
            document.add_heading(f"{e.title} — {e.company}", level=2)
            meta = join_parts(e.location, date_range(e.start_date, e.end_date))
            if meta:
                document.add_paragraph(meta)
            bullets(e.bullets)
    if resume.projects:
        document.add_heading("Projects", level=1)
        for p in resume.projects:
            tech = f" ({', '.join(p.technologies)})" if p.technologies else ""
            document.add_heading(f"{p.name}{tech}", level=2)
            meta = join_parts(p.description, p.link)
            if meta:
                document.add_paragraph(meta)
            bullets(p.bullets)
    if resume.education:
        document.add_heading("Education", level=1)
        for ed in resume.education:
            document.add_heading(ed.institution, level=2)
            meta = join_parts(
                join_parts(ed.degree, ed.field_of_study, sep=", "), ed.location,
                date_range(ed.start_date, ed.end_date), ed.grade,
            )
            if meta:
                document.add_paragraph(meta)
            bullets(ed.highlights)
    for title, items in (
        ("Certifications", resume.certifications),
        ("Achievements", resume.achievements),
        ("Languages", resume.languages),
    ):
        if items:
            document.add_heading(title, level=1)
            bullets(items)
    return _docx_bytes(document)


def cover_letter_docx(letter: CoverLetter, contact: ContactInfo) -> bytes:
    document = _new_document(11)
    if contact.name:
        document.add_paragraph().add_run(contact.name).bold = True
    contact_line = _contact_lines(contact)[0]
    if contact_line:
        document.add_paragraph(contact_line)
    document.add_paragraph(letter.greeting)
    for paragraph in letter.paragraphs:
        document.add_paragraph(paragraph)
    document.add_paragraph(letter.sign_off)
    if letter.signature:
        document.add_paragraph(letter.signature)
    return _docx_bytes(document)
