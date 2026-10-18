import io

import docx
import pymupdf

from app.generators.export import (
    cover_letter_docx,
    cover_letter_pdf,
    resume_docx,
    resume_pdf,
    safe_filename,
)
from app.llm.models import ContactInfo, CoverLetter, Education, Experience, ParsedResume, Project

RESUME = ParsedResume(
    contact=ContactInfo(name="Jane Doe", email="jane@example.com", linkedin="https://linkedin.com/in/jane"),
    summary="Backend engineer.",
    skills=["Python", "FastAPI"],
    experience=[Experience(
        title="Engineer", company="Acme", start_date="Jan 2023", end_date="Present",
        bullets=["Built APIs serving 2M requests/day with p95 < 80 ms & zero downtime"],
    )],
    projects=[Project(name="AutoApply", technologies=["LangGraph"], bullets=["Built an agent"])],
    education=[Education(institution="UT Austin", degree="B.S.", location="Austin, TX")],
)
LETTER = CoverLetter(
    greeting="Dear Hiring Manager,", paragraphs=["First — with a dash.", "Second."],
    sign_off="Sincerely,", signature="Jane Doe",
)


def pdf_text(data: bytes) -> str:
    doc = pymupdf.open(stream=data, filetype="pdf")
    return "".join(doc[i].get_text() for i in range(doc.page_count))


def test_resume_pdf_has_selectable_text_with_special_characters():
    data = resume_pdf(RESUME)
    assert data.startswith(b"%PDF")
    text = pdf_text(data)
    for expected in ("Jane Doe", "Experience", "Engineer — Acme", "p95 < 80 ms & zero downtime",
                     "AutoApply (LangGraph)", "B.S. | Austin, TX"):
        assert expected in text


def test_resume_pdf_fonts_are_subset():
    assert len(resume_pdf(RESUME)) < 100_000


def test_long_resume_spans_pages():
    long = RESUME.model_copy(deep=True)
    long.experience[0].bullets = [f"Achievement number {i} described in a full sentence" for i in range(120)]
    doc = pymupdf.open(stream=resume_pdf(long), filetype="pdf")
    assert doc.page_count > 1
    assert "Achievement number 119" in pdf_text(resume_pdf(long))


def test_cover_letter_pdf():
    text = pdf_text(cover_letter_pdf(LETTER, RESUME.contact))
    assert "Dear Hiring Manager," in text and "First — with a dash." in text
    assert text.index("Sincerely,") < text.rindex("Jane Doe")


def test_resume_docx_structure():
    document = docx.Document(io.BytesIO(resume_docx(RESUME)))
    texts = [p.text for p in document.paragraphs]
    assert texts[0] == "Jane Doe"
    assert "Engineer — Acme" in texts
    bullet = next(p for p in document.paragraphs if p.text.startswith("Built APIs"))
    assert bullet.style is not None and bullet.style.name == "List Bullet"


def test_cover_letter_docx():
    document = docx.Document(io.BytesIO(cover_letter_docx(LETTER, RESUME.contact)))
    assert [p.text for p in document.paragraphs][-4:] == [
        "First — with a dash.", "Second.", "Sincerely,", "Jane Doe"
    ]


def test_safe_filename():
    assert safe_filename("Jane O'Doe", "Résumé") == "Jane_O_Doe_R_sum"
    assert safe_filename("", "  ") == "document"


def pdf_links(data: bytes) -> list[str]:
    doc = pymupdf.open(stream=data, filetype="pdf")
    return [link["uri"] for i in range(doc.page_count) for link in doc[i].get_links()]


def test_resume_pdf_links_match_visible_text():
    resume = RESUME.model_copy(deep=True)
    resume.contact.github = "github.com/jane"
    resume.projects[0].link = "github.com/jane/autoapply"
    links = pdf_links(resume_pdf(resume))
    assert links.count("mailto:jane@example.com") == 1
    assert "https://linkedin.com/in/jane" in links
    assert links.count("https://github.com/jane") == 1  # not also placed over the project link
    assert links.count("https://github.com/jane/autoapply") == 1


def test_cover_letter_pdf_links_email():
    assert pdf_links(cover_letter_pdf(LETTER, RESUME.contact)) == ["mailto:jane@example.com"]
