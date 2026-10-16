import io
from unittest.mock import patch

import docx
import pymupdf
import pytest
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from app.config import settings
from app.llm.models import ContactInfo, Education, Experience, ParsedResume
from app.llm.provider import StructuredOutputError
from app.parsers.resume_parser import (
    FileType,
    ResumeParseError,
    _normalize,
    detect_file_type,
    extract_text,
    load_resume_file,
    parse_resume,
    structure_resume,
)

RESUME_TEXT = """Jane Doe
jane.doe@example.com | +1 555 0100 | Austin, TX

EXPERIENCE
Software Engineer, Acme Corp (Jan 2022 - Present)
- Built REST APIs in Python and FastAPI serving 2M requests per day
- Reduced deployment time by 40% with Docker and GitHub Actions

EDUCATION
B.S. Computer Science, University of Texas (2018 - 2022)
"""

SAMPLE_PARSED = ParsedResume(
    contact=ContactInfo(name="Jane Doe", email="jane.doe@example.com"),
    skills=["Python", "python", "FastAPI", " Docker ", ""],
    experience=[Experience(title="Software Engineer", company="Acme Corp")],
    education=[Education(institution="University of Texas")],
)


def make_pdf(text: str = RESUME_TEXT, link: str | None = None, **save_kwargs) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    if text:
        page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), text, fontsize=10)
    if link:
        page.insert_link(
            {"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(50, 50, 200, 70), "uri": link}
        )
    data = doc.tobytes(**save_kwargs)
    doc.close()
    return data


def make_docx() -> bytes:
    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "Jane Doe | jane.doe@example.com"
    for line in RESUME_TEXT.splitlines()[2:]:
        document.add_paragraph(line)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Kubernetes"
    table.cell(0, 1).text = "PostgreSQL"
    document.part.relate_to("https://github.com/janedoe", RT.HYPERLINK, is_external=True)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# --- File type detection ---


def test_detects_pdf_by_content_not_extension():
    assert detect_file_type(make_pdf(), "resume.docx") is FileType.PDF


def test_detects_docx():
    assert detect_file_type(make_docx(), "resume") is FileType.DOCX


def test_detects_txt_by_extension():
    assert detect_file_type(b"plain text", "resume.txt") is FileType.TXT


def test_rejects_legacy_doc():
    with pytest.raises(ResumeParseError, match="Legacy .doc"):
        detect_file_type(b"\xd0\xcf\x11\xe0 binary", "resume.doc")


def test_rejects_non_word_zip():
    buf = io.BytesIO()
    import zipfile

    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("hello.txt", "hi")
    with pytest.raises(ResumeParseError, match="Unsupported"):
        detect_file_type(buf.getvalue(), "resume.zip")


def test_rejects_unknown_type():
    with pytest.raises(ResumeParseError, match="Unsupported"):
        detect_file_type(b"\x89PNG\r\n", "photo.png")


# --- Text extraction ---


def test_extracts_pdf_text_and_links():
    text = extract_text(make_pdf(link="https://linkedin.com/in/janedoe"), "r.pdf")
    assert "Jane Doe" in text
    assert "FastAPI" in text
    assert "Links found in document:\nhttps://linkedin.com/in/janedoe" in text


def test_extracts_docx_header_body_table_and_links():
    text = extract_text(make_docx(), "r.docx")
    assert text.index("Jane Doe") < text.index("EXPERIENCE")
    assert "Kubernetes" in text and "PostgreSQL" in text
    assert "https://github.com/janedoe" in text


def test_extracts_txt():
    assert "Acme Corp" in extract_text(RESUME_TEXT.encode(), "r.txt")


def test_rejects_empty_file():
    with pytest.raises(ResumeParseError, match="empty"):
        extract_text(b"", "r.pdf")


def test_rejects_oversized_file(monkeypatch):
    monkeypatch.setattr(settings, "max_resume_size_mb", 1)
    with pytest.raises(ResumeParseError, match="1 MB limit"):
        extract_text(b"%PDF" + b"0" * (1024 * 1024), "r.pdf")


def test_rejects_scanned_pdf():
    with pytest.raises(ResumeParseError, match="scanned image"):
        extract_text(make_pdf(text=""), "r.pdf")


def test_rejects_encrypted_pdf():
    # PyMuPDF's stubs omit its encryption constants.
    aes_256 = pymupdf.PDF_ENCRYPT_AES_256  # type: ignore[attr-defined]
    data = make_pdf(encryption=aes_256, owner_pw="o", user_pw="u")
    with pytest.raises(ResumeParseError, match="password-protected"):
        extract_text(data, "r.pdf")


def test_rejects_corrupted_pdf():
    with pytest.raises(ResumeParseError):
        extract_text(b"%PDF-1.7 not really a pdf", "r.pdf")


def test_rejects_too_little_text():
    with pytest.raises(ResumeParseError, match="too little text"):
        extract_text(b"Jane Doe", "r.txt")


def test_normalize_fixes_ligatures_bullets_and_blank_lines():
    assert _normalize("ﬁnance Python  \n\n\n\nDocker") == "finance• Python\n\nDocker"


def test_load_resume_file_missing():
    with pytest.raises(ResumeParseError, match="not found"):
        load_resume_file("does/not/exist.pdf")


# --- LLM structuring ---


def parsed_sample() -> ParsedResume:
    return SAMPLE_PARSED.model_copy(deep=True)


@patch("app.parsers.resume_parser.invoke_structured")
def test_structure_resume_dedupes_skills(mock_invoke):
    mock_invoke.return_value = parsed_sample()
    result = structure_resume(RESUME_TEXT)
    assert mock_invoke.call_args[0][0] is ParsedResume
    assert result.skills == ["Python", "FastAPI", "Docker"]
    assert result.contact.name == "Jane Doe"


@patch("app.parsers.resume_parser.invoke_structured")
def test_structure_resume_wraps_text_as_data(mock_invoke):
    mock_invoke.return_value = parsed_sample()
    structure_resume(RESUME_TEXT)
    assert f"<resume>\n{RESUME_TEXT}\n</resume>" in mock_invoke.call_args[0][1]


@patch("app.parsers.resume_parser.invoke_structured")
def test_structure_resume_truncates_long_text(mock_invoke):
    mock_invoke.return_value = parsed_sample()
    structure_resume("x" * 50_000)
    assert "x" * 30_001 not in mock_invoke.call_args[0][1]


@patch("app.parsers.resume_parser.invoke_structured")
def test_structure_resume_rejects_non_resume(mock_invoke):
    mock_invoke.return_value = ParsedResume(skills=["cooking"])
    with pytest.raises(ResumeParseError, match="doesn't look like a resume"):
        structure_resume("A recipe for banana bread " * 10)


@patch(
    "app.parsers.resume_parser.invoke_structured",
    side_effect=StructuredOutputError("no valid ParsedResume"),
)
def test_structure_resume_handles_invalid_llm_output(_mock_invoke):
    with pytest.raises(ResumeParseError, match="try again"):
        structure_resume(RESUME_TEXT)


@patch("app.parsers.resume_parser.invoke_structured")
def test_parse_resume_from_bytes_and_path(mock_invoke, tmp_path):
    mock_invoke.side_effect = lambda *_: parsed_sample()
    assert parse_resume(make_pdf(), "r.pdf").contact.name == "Jane Doe"

    path = tmp_path / "resume.docx"
    path.write_bytes(make_docx())
    assert parse_resume(path).contact.name == "Jane Doe"
